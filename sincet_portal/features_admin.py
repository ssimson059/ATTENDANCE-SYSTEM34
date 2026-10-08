"""Admin (and faculty report) side of the v7 features. Routes are attached to features.bp."""
import csv
import io
import os
import re
from datetime import date

from flask import Response, abort, jsonify, render_template, request, session

import db
import notify
from features import (A, ADMIN_SCOPES, DAYS, PERIOD_TIMES, SESSIONS, _audit, _faculty, _int, _json, _num, _date,
                      bp, contacts_map, gateway_enabled, need, open_fees, pdf_response, reminder_channels,
                      run_reminders, send_parent_alerts, student_results, build_marksheet_pdf, _announce_fields)


def _scope():
    return session.get("scope") or "super"


# =========================================================================== REPORTS
def low_attendance(threshold, dept=None, year=None, section=None, min_classes=1):
    sql = ("SELECT u.id AS student_id, u.username, u.name, d.code AS dept, u.year_no, s.name AS section, "
           "SUM(CASE WHEN a.status='Present' THEN 1 ELSE 0 END) AS present, COUNT(a.id) AS total "
           "FROM users u JOIN departments d ON d.id=u.department_id LEFT JOIN sections s ON s.id=u.section_id "
           "JOIN attendance a ON a.student_id=u.id WHERE u.role='student' AND u.status='active'")
    args = []
    for col, val in (("u.department_id", dept), ("u.year_no", year), ("u.section_id", section)):
        if val:
            sql += f" AND {col}=%s"
            args.append(val)
    sql += " GROUP BY u.id, u.username, u.name, d.code, u.year_no, s.name HAVING COUNT(a.id)>=%s"
    out = []
    for r in db.query(sql, tuple(args + [min_classes])):
        pct = round(100.0 * (r["present"] or 0) / r["total"], 1)
        if pct < threshold:
            r.update(pct=pct, threshold=threshold, present=int(r["present"] or 0), total=int(r["total"]))
            out.append(r)
    return sorted(out, key=lambda r: (r["pct"], r["username"]))


def fee_summary():
    depts = db.query("SELECT id, code, name FROM departments ORDER BY code")
    stud = {r["department_id"]: r["n"] for r in db.query(
        "SELECT department_id, COUNT(*) AS n FROM users WHERE role='student' AND status='active' GROUP BY department_id")}
    billed = {r["department_id"]: float(r["t"]) for r in db.query(
        "SELECT u.department_id, COALESCE(SUM(f.amount),0) AS t FROM fees f JOIN users u ON u.id=f.student_id "
        "GROUP BY u.department_id")}
    pay = {r["department_id"]: r for r in db.query(
        "SELECT u.department_id, COALESCE(SUM(CASE WHEN p.status='received' THEN p.amount ELSE 0 END),0) AS rec, "
        "COALESCE(SUM(CASE WHEN p.status='pending' THEN p.amount ELSE 0 END),0) AS pen "
        "FROM fee_payments p JOIN users u ON u.id=p.student_id GROUP BY u.department_id")}
    rows = []
    for d in depts:
        b, p = billed.get(d["id"], 0.0), pay.get(d["id"], {})
        rec, pen = float(p.get("rec", 0) or 0), float(p.get("pen", 0) or 0)
        rows.append(dict(dept=d["code"], name=d["name"], students=stud.get(d["id"], 0), billed=round(b, 2),
                         received=round(rec, 2), pending=round(pen, 2), outstanding=round(max(0.0, b - rec), 2),
                         pct=round(100 * rec / b, 1) if b else 0.0))
    return rows


def pending_fees(dept=None, year=None):
    rows = [r for r in open_fees(dept, year) if r["due_amount"] > 0]
    return sorted(rows, key=lambda r: (r["dept_code"], r["username"]))


def _report_data(name, a):
    dept, year = _int(a.get("dept")) if a.get("dept") else None, _int(a.get("year")) if a.get("year") else None
    if name == "fee_summary":
        rows = fee_summary()
        return ["Department", "Name", "Students", "Billed", "Received", "Awaiting verification", "Outstanding", "% collected"], \
            [[r["dept"], r["name"], r["students"], r["billed"], r["received"], r["pending"], r["outstanding"], r["pct"]] for r in rows], rows
    if name == "pending_fees":
        rows = pending_fees(dept, year)
        return ["Student ID", "Name", "Dept", "Year", "Fee", "Amount", "Paid", "Due", "Awaiting verification", "Due date"], \
            [[r["username"], r["name"], r["dept_code"], r["year_no"], r["title"], r["amount"], r["paid"], r["due_amount"],
              r["pending"], r["due_date"] or ""] for r in rows], rows
    if name == "low_attendance":
        thr = _num(a.get("threshold", A.ATT_THRESHOLD), 1, 100) or A.ATT_THRESHOLD
        rows = low_attendance(thr, dept, year, _int(a.get("section")) if a.get("section") else None)
        return ["Student ID", "Name", "Dept", "Year", "Section", "Present", "Total classes", "Attendance %"], \
            [[r["username"], r["name"], r["dept"], r["year_no"], r["section"] or "", r["present"], r["total"], r["pct"]]
             for r in rows], rows
    return None, None, None


_REPORT_SCOPES = {"fee_summary": ("super", "accounts"), "pending_fees": ("super", "accounts"),
                  "low_attendance": ("super", "exam")}
_TITLES = {"fee_summary": "Fee collection - department-wise", "pending_fees": "Pending fees list",
           "low_attendance": "Low attendance list"}


def send_xlsx(name, title, header, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return A.send_csv(name, header, rows)
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append(header)
    for r in rows:
        ws.append([A._safe_cell(c) if isinstance(c, str) else c for c in r])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B2A5B")
        c.alignment = Alignment(horizontal="center")
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(40, max(len(str(c.value or "")) for c in col) + 3)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={name}.xlsx"})


def export_table(fmt, name, title, header, rows):
    if fmt == "csv":
        return A.send_csv(name, header, rows)
    if fmt == "xlsx":
        return send_xlsx(name, title, header, rows)
    if fmt == "pdf":
        return A.send_pdf(name, title, header, rows)
    abort(404)


@bp.route("/admin/reports")
@need("admin")
def admin_reports_page():
    return render_template("admin/reports.html", active="reports", scope=_scope(), allowed=[
        k for k, v in _REPORT_SCOPES.items() if _scope() in v],
        depts=db.query("SELECT id, code, name FROM departments ORDER BY name"), titles=_TITLES,
        default_threshold=A.ATT_THRESHOLD)


@bp.route("/admin/api/reports/<name>")
@need("admin")
def admin_report_api(name):
    if name not in _REPORT_SCOPES or _scope() not in _REPORT_SCOPES[name]:
        return A.bad("Your admin role cannot open this report.", 403)
    header, body, _rows = _report_data(name, request.args)
    return jsonify(ok=True, header=header, rows=body)


@bp.route("/admin/reports/export/<name>.<fmt>")
@need("admin")
def admin_report_export(name, fmt):
    if name not in _REPORT_SCOPES or _scope() not in _REPORT_SCOPES[name]:
        abort(403)
    header, body, _rows = _report_data(name, request.args)
    A.audit("report_export", f"{name} {fmt}")
    return export_table(fmt, name, _TITLES[name], header, body)


@bp.route("/admin/api/reports/alert-parents", methods=["POST"])
@need("admin")
def admin_alert_parents():
    d = _json()
    kind, channel = d.get("kind"), d.get("channel", "sms")
    if kind not in ("low_attendance", "fee_pending") or channel not in ("sms", "whatsapp", "email"):
        return A.bad("Invalid request.")
    need_scope = ("super", "exam") if kind == "low_attendance" else ("super", "accounts")
    if _scope() not in need_scope:
        return A.bad("Your admin role cannot send this alert.", 403)
    args = dict(dept=d.get("dept"), year=d.get("year"), threshold=d.get("threshold", A.ATT_THRESHOLD))
    _h, _b, rows = _report_data("low_attendance" if kind == "low_attendance" else "pending_fees", args)
    if d.get("preview"):
        cm = contacts_map()
        key = "parent_email" if channel == "email" else "parent_phone"
        return jsonify(ok=True, students=len(rows), with_contact=sum(1 for r in rows if (cm.get(r["student_id"]) or {}).get(key)),
                       configured=notify.configured(channel))
    res = send_parent_alerts(kind, rows, channel)
    A.audit("parent_alerts", f"{kind} {channel} {res}")
    return jsonify(ok=True, **res, configured=notify.configured(channel))


# ----- faculty: low attendance of own department
@bp.route("/faculty/reports")
@need("faculty")
def faculty_reports_page():
    me = _faculty()
    return render_template("faculty/report.html", me=me, default_threshold=A.ATT_THRESHOLD, sections=db.query(
        "SELECT id, year_no, name FROM sections WHERE department_id=%s ORDER BY year_no, name", (me["department_id"],)))


def _faculty_args(me):
    a = dict(request.args)
    a["dept"] = me["department_id"]
    return a


@bp.route("/faculty/api/reports/low-attendance")
@need("faculty")
def faculty_low_att():
    header, body, _r = _report_data("low_attendance", _faculty_args(_faculty()))
    return jsonify(ok=True, header=header, rows=body)


@bp.route("/faculty/reports/low-attendance.<fmt>")
@need("faculty")
def faculty_low_att_export(fmt):
    header, body, _r = _report_data("low_attendance", _faculty_args(_faculty()))
    return export_table(fmt, "low_attendance", "Low attendance list", header, body)


@bp.route("/faculty/api/reports/alert-parents", methods=["POST"])
@need("faculty")
def faculty_alert_parents():
    d = _json()
    channel = d.get("channel", "sms")
    if channel not in ("sms", "whatsapp", "email"):
        return A.bad("Invalid channel.")
    me = _faculty()
    args = dict(dept=me["department_id"], year=d.get("year"), section=d.get("section"),
                threshold=d.get("threshold", A.ATT_THRESHOLD))
    _h, _b, rows = _report_data("low_attendance", args)
    if d.get("preview"):
        cm = contacts_map()
        key = "parent_email" if channel == "email" else "parent_phone"
        return jsonify(ok=True, students=len(rows), with_contact=sum(1 for r in rows if (cm.get(r["student_id"]) or {}).get(key)),
                       configured=notify.configured(channel))
    res = send_parent_alerts("low_attendance", rows, channel)
    A.audit("parent_alerts", f"low_attendance {channel} {res}")
    return jsonify(ok=True, **res, configured=notify.configured(channel))


# =========================================================================== NOTIFICATIONS (contacts, outbox, reminders)
@bp.route("/admin/notifications")
@need("admin")
def admin_notifications_page():
    return render_template("admin/notifications.html", active="notify", depts=db.query(
        "SELECT id, code, name FROM departments ORDER BY name"), channels=notify.status(),
        reminder_days=os.getenv("REMINDER_DAYS", "7,3,1,0"), reminder_channels=reminder_channels(),
        gateway=gateway_enabled())


_PHONE = re.compile(r"^\+?[0-9 \-]{10,15}$")
_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _clean_contact(d):
    out = {}
    for k in ("phone", "parent_phone"):
        v = str(d.get(k, "") or "").strip()
        if v and not _PHONE.match(v):
            return None, f"Invalid phone number: {v}"
        out[k] = v or None
    for k in ("email", "parent_email"):
        v = str(d.get(k, "") or "").strip()
        if v and (not _EMAIL.match(v) or len(v) > 120):
            return None, f"Invalid email: {v}"
        out[k] = v or None
    return out, None


def _save_contact(sid, c):
    if db.query_one("SELECT student_id FROM student_contacts WHERE student_id=%s", (sid,)):
        db.execute("UPDATE student_contacts SET phone=%s, email=%s, parent_phone=%s, parent_email=%s, updated_at=%s "
                   "WHERE student_id=%s", (c["phone"], c["email"], c["parent_phone"], c["parent_email"], A.now(), sid),
                   commit=False)
    else:
        db.execute("INSERT INTO student_contacts (student_id, phone, email, parent_phone, parent_email, updated_at) "
                   "VALUES (%s,%s,%s,%s,%s,%s)", (sid, c["phone"], c["email"], c["parent_phone"], c["parent_email"],
                                                   A.now()), commit=False)


@bp.route("/admin/api/notify/contacts")
@need("admin")
def api_contacts():
    q, dept = request.args.get("q", "").strip(), _int(request.args.get("dept")) if request.args.get("dept") else None
    sql = ("SELECT u.id, u.username, u.name, d.code AS dept, u.year_no, c.phone, c.email, c.parent_phone, c.parent_email "
           "FROM users u JOIN departments d ON d.id=u.department_id LEFT JOIN student_contacts c ON c.student_id=u.id "
           "WHERE u.role='student' AND u.status='active'")
    args = []
    if q:
        sql += " AND (u.username LIKE %s OR u.name LIKE %s)"
        args += [f"%{q}%"] * 2
    if dept:
        sql += " AND u.department_id=%s"
        args.append(dept)
    return jsonify(ok=True, rows=db.query(sql + " ORDER BY u.username LIMIT 300", tuple(args)))


@bp.route("/admin/api/notify/contacts", methods=["POST"])
@need("admin")
def api_contact_save():
    d = _json()
    sid = _int(d.get("student_id"))
    if sid is None or not db.query_one("SELECT id FROM users WHERE id=%s AND role='student'", (sid,)):
        return A.bad("Student not found.", 404)
    c, err = _clean_contact(d)
    if err:
        return A.bad(err)
    with db.transaction():
        _save_contact(sid, c)
    A.audit("contact_saved", f"student {sid}")
    return jsonify(ok=True)


@bp.route("/admin/api/notify/contacts/import", methods=["POST"])
@need("admin")
def api_contact_import():
    """CSV text: username,phone,email,parent_phone,parent_email (header row optional)."""
    text = str(_json().get("csv", ""))
    if len(text) > 400000:
        return A.bad("File too large.")
    done, skipped = 0, []
    users = {r["username"]: r["id"] for r in db.query("SELECT id, username FROM users WHERE role='student'")}
    with db.transaction():
        for row in csv.reader(io.StringIO(text)):
            if not row or row[0].strip().lower() in ("username", "student", "student id", ""):
                continue
            row = [x.strip() for x in row] + [""] * 5
            sid = users.get(row[0].upper())
            c, err = _clean_contact(dict(phone=row[1], email=row[2], parent_phone=row[3], parent_email=row[4]))
            if not sid or err:
                skipped.append(row[0])
                continue
            _save_contact(sid, c)
            done += 1
    A.audit("contacts_imported", f"{done} saved, {len(skipped)} skipped")
    return jsonify(ok=True, saved=done, skipped=skipped[:20], skipped_count=len(skipped))


@bp.route("/admin/api/notify/outbox")
@need("admin")
def api_outbox():
    st = request.args.get("status", "")
    sql, args = "SELECT id, channel, recipient, subject, body, kind, status, error, created_at FROM outbox", []
    if st in ("sent", "logged", "failed"):
        sql += " WHERE status=%s"
        args.append(st)
    return jsonify(ok=True, rows=db.query(sql + " ORDER BY id DESC LIMIT 200", tuple(args)))


@bp.route("/admin/api/notify/reminders/run", methods=["POST"])
@need("admin")
def api_run_reminders():
    return jsonify(ok=True, **run_reminders(session.get("username", "admin")), channels=notify.status())


@bp.route("/internal/cron/reminders")
def cron_reminders():
    """For an external scheduler:  GET /internal/cron/reminders?token=<CRON_TOKEN>  (once a day)."""
    import hmac
    token = os.getenv("CRON_TOKEN", "")
    if not token or not hmac.compare_digest(token, request.args.get("token", "")):
        abort(404)
    return jsonify(ok=True, **run_reminders("cron"))


# =========================================================================== SCHEDULE (timetable + exams)
@bp.route("/admin/schedule")
@need("admin")
def admin_schedule_page():
    return render_template("admin/schedule.html", active="schedule", exams=A.EXAMS, sessions=SESSIONS)


@bp.route("/admin/api/schedule/meta")
@need("admin")
def api_schedule_meta():
    return jsonify(
        ok=True, periods=PERIOD_TIMES, days=DAYS,
        departments=db.query("SELECT id, code, name FROM departments ORDER BY code"),
        sections=db.query("SELECT s.id, s.department_id, s.year_no, s.name, d.code AS dept FROM sections s "
                          "JOIN departments d ON d.id=s.department_id ORDER BY d.code, s.year_no, s.name"),
        subjects=db.query("SELECT id, department_id, year_no, semester, code, name FROM subjects ORDER BY code"))


@bp.route("/admin/api/schedule/timetable", methods=["GET", "POST"])
@need("admin")
def api_timetable():
    if request.method == "GET":
        sec = _int(request.args.get("section"))
        if sec is None:
            return A.bad("Select a section.")
        return jsonify(ok=True, cells=db.query(
            "SELECT day_no, period_no, subject_id, label, room FROM timetable WHERE section_id=%s", (sec,)))
    d = _json()
    sec = _int(d.get("section"))
    s = db.query_one("SELECT id, department_id, year_no FROM sections WHERE id=%s", (sec,)) if sec else None
    if not s or not isinstance(d.get("cells"), list):
        return A.bad("Select a section.")
    rows = []
    for c in d["cells"]:
        day, per = _int(c.get("day")), _int(c.get("period"))
        sub = _int(c.get("subject_id")) if str(c.get("subject_id", "")).strip() else None
        label, room = str(c.get("label", "")).strip()[:60] or None, str(c.get("room", "")).strip()[:30] or None
        if not (day and per and 1 <= day <= len(DAYS) and 1 <= per <= len(PERIOD_TIMES)):
            return A.bad("Invalid day / period.")
        if not (sub or label):
            continue
        if sub and not db.query_one("SELECT id FROM subjects WHERE id=%s AND department_id=%s AND year_no=%s",
                                    (sub, s["department_id"], s["year_no"])):
            return A.bad("A selected subject does not belong to this department / year.")
        rows.append((s["department_id"], s["year_no"], sec, day, per, sub, label, room))
    with db.transaction():
        db.execute("DELETE FROM timetable WHERE section_id=%s", (sec,), commit=False)
        for r in rows:
            db.execute("INSERT INTO timetable (department_id, year_no, section_id, day_no, period_no, subject_id, label, "
                       "room) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)", r, commit=False)
    A.audit("timetable_saved", f"section {sec} {len(rows)} periods")
    return jsonify(ok=True, saved=len(rows))


@bp.route("/admin/api/schedule/exams", methods=["GET", "POST"])
@need("admin")
def api_exams():
    if request.method == "GET":
        sql = ("SELECT e.id, e.exam_type, e.exam_date, e.session, e.venue, e.year_no, d.code AS dept, sb.code AS subject_code, "
               "sb.name AS subject_name FROM exam_schedule e JOIN subjects sb ON sb.id=e.subject_id "
               "JOIN departments d ON d.id=e.department_id WHERE 1=1")
        args = []
        for col, key in (("e.department_id", "dept"), ("e.year_no", "year")):
            if request.args.get(key, "").isdigit():
                sql += f" AND {col}=%s"
                args.append(int(request.args[key]))
        return jsonify(ok=True, rows=db.query(sql + " ORDER BY e.exam_type, e.exam_date, e.session LIMIT 500", tuple(args)))
    d = _json()
    sub = db.query_one("SELECT id, department_id, year_no FROM subjects WHERE id=%s", (_int(d.get("subject_id")),))
    when = _date(d.get("exam_date"))
    if not sub or d.get("exam_type") not in A.EXAMS or not when or d.get("session") not in SESSIONS:
        return A.bad("Choose subject, exam, a valid date and FN / AN.")
    venue = str(d.get("venue", "")).strip()[:60] or None
    old = db.query_one("SELECT id FROM exam_schedule WHERE subject_id=%s AND exam_type=%s", (sub["id"], d["exam_type"]))
    if old:
        db.execute("UPDATE exam_schedule SET exam_date=%s, session=%s, venue=%s WHERE id=%s", (when, d["session"], venue, old["id"]))
    else:
        db.insert("INSERT INTO exam_schedule (department_id, year_no, subject_id, exam_type, exam_date, session, venue) "
                  "VALUES (%s,%s,%s,%s,%s,%s,%s)", (sub["department_id"], sub["year_no"], sub["id"], d["exam_type"], when,
                                                    d["session"], venue))
    A.audit("exam_scheduled", f"subject {sub['id']} {d['exam_type']} {when}")
    return jsonify(ok=True)


@bp.route("/admin/api/schedule/exams/<int:eid>", methods=["DELETE"])
@need("admin")
def api_exam_delete(eid):
    if not db.execute("DELETE FROM exam_schedule WHERE id=%s", (eid,)):
        return A.bad("Not found.", 404)
    A.audit("exam_unscheduled", str(eid))
    return jsonify(ok=True)


# =========================================================================== RESULTS (admin lookup + mark sheet)
@bp.route("/admin/results")
@need("admin")
def admin_results_page():
    return render_template("admin/results.html", active="results")


def _stu_by_username(u):
    return db.query_one("SELECT u.id, u.name, u.username, u.year_no, d.name AS dept_name FROM users u JOIN departments d "
                        "ON d.id=u.department_id WHERE u.role='student' AND u.username=%s", (u.strip().upper(),))


@bp.route("/admin/api/results")
@need("admin")
def admin_results_api():
    s = _stu_by_username(request.args.get("username", ""))
    if not s:
        return A.bad("Student not found.", 404)
    return jsonify(ok=True, student=dict(name=s["name"], username=s["username"], dept=s["dept_name"], year=s["year_no"]),
                   res=student_results(s["id"]))


@bp.route("/admin/results/marksheet.pdf")
@need("admin")
def admin_marksheet():
    s = _stu_by_username(request.args.get("username", ""))
    if not s:
        abort(404)
    sem = _int(request.args.get("sem")) if request.args.get("sem", "") != "" else None
    return pdf_response(f"marksheet_{s['username']}.pdf", build_marksheet_pdf(s, student_results(s["id"]), sem))


# =========================================================================== ANNOUNCEMENTS (admin)
@bp.route("/admin/announcements")
@need("admin")
def admin_announcements_page():
    return render_template("admin/announcements.html", active="announce", depts=db.query(
        "SELECT id, code, name FROM departments ORDER BY name"))


@bp.route("/admin/api/announcements", methods=["GET", "POST"])
@need("admin")
def api_announcements():
    if request.method == "GET":
        return jsonify(ok=True, rows=db.query(
            "SELECT a.id, a.title, a.body, a.audience, a.pinned, a.author_name, a.author_role, a.year_no, a.created_at, "
            "a.expires_on, d.code AS dept FROM announcements a LEFT JOIN departments d ON d.id=a.department_id "
            "ORDER BY a.id DESC LIMIT 200"))
    d = _json()
    f, err = _announce_fields(d)
    if err:
        return A.bad(err)
    aud = d.get("audience", "all")
    dept = _int(d.get("dept")) if str(d.get("dept", "")).strip() else None
    year = _int(d.get("year")) if str(d.get("year", "")).strip() else None
    if aud not in ("all", "student", "faculty"):
        return A.bad("Invalid audience.")
    aid = db.insert("INSERT INTO announcements (title, body, audience, department_id, year_no, pinned, author, author_name, "
                    "author_role, created_at, expires_on) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,'admin',%s,%s)",
                    (f["title"], f["body"], aud, dept, year, 1 if d.get("pinned") else 0, session["username"],
                     session.get("name") or "Admin", A.now(), f["expires_on"]))
    A.audit("announcement_created", f"{aid} {f['title']}")
    return jsonify(ok=True, id=aid)


@bp.route("/admin/api/announcements/<int:aid>", methods=["DELETE"])
@need("admin")
def api_announcement_delete(aid):
    if not db.execute("DELETE FROM announcements WHERE id=%s", (aid,)):
        return A.bad("Not found.", 404)
    A.audit("announcement_deleted", str(aid))
    return jsonify(ok=True)


# =========================================================================== STAFF (role-wise admin accounts) - super only
@bp.route("/admin/staff")
@need("admin")
def admin_staff_page():
    return render_template("admin/staff.html", active="staff", scopes=ADMIN_SCOPES)


@bp.route("/admin/api/staff")
@need("admin")
def api_staff():
    return jsonify(ok=True, rows=db.query("SELECT id, username, full_name, scope, last_login, created_at FROM admins "
                                          "ORDER BY id"), scopes=ADMIN_SCOPES)


@bp.route("/admin/api/staff", methods=["POST"])
@need("admin")
def api_staff_create():
    d = _json()
    uname, name, scope = str(d.get("username", "")).strip(), str(d.get("name", "")).strip(), d.get("scope")
    if not re.fullmatch(r"[A-Za-z0-9_.\-]{3,40}", uname) or not (2 <= len(name) <= 100) or scope not in ADMIN_SCOPES:
        return A.bad("Enter a username (3-40 letters/digits), a name and choose a role.")
    pw = A.gen_password(10)
    try:
        db.insert("INSERT INTO admins (username, password_hash, full_name, scope, created_at) VALUES (%s,%s,%s,%s,%s)",
                  (uname, A.hash_pw(pw), name, scope, A.now()))
    except db.IntegrityViolation:
        return A.bad("That username already exists.", 409)
    A.audit("staff_created", f"{uname} ({scope})")
    return jsonify(ok=True, username=uname, password=pw)


@bp.route("/admin/api/staff/<int:aid>/reset", methods=["POST"])
@need("admin")
def api_staff_reset(aid):
    pw = A.gen_password(10)
    n = db.execute("UPDATE admins SET password_hash=%s, failed_attempts=0, locked_until=NULL, last_login=NULL WHERE id=%s",
                   (A.hash_pw(pw), aid))
    if not n:
        return A.bad("Not found.", 404)
    A.audit("staff_password_reset", str(aid))
    return jsonify(ok=True, password=pw)


@bp.route("/admin/api/staff/<int:aid>", methods=["DELETE"])
@need("admin")
def api_staff_delete(aid):
    row = db.query_one("SELECT id, username, scope FROM admins WHERE id=%s", (aid,))
    if not row:
        return A.bad("Not found.", 404)
    if row["id"] == session["uid"]:
        return A.bad("You cannot delete your own account.", 409)
    if row["scope"] == "super" and db.query_one("SELECT COUNT(*) AS n FROM admins WHERE scope='super'")["n"] <= 1:
        return A.bad("At least one super admin must remain.", 409)
    db.execute("DELETE FROM admins WHERE id=%s", (aid,))
    A.audit("staff_deleted", row["username"])
    return jsonify(ok=True)


@bp.route("/admin/api/staff/<int:aid>/scope", methods=["POST"])
@need("admin")
def api_staff_scope(aid):
    scope = _json().get("scope")
    row = db.query_one("SELECT id, scope FROM admins WHERE id=%s", (aid,))
    if scope not in ADMIN_SCOPES or not row:
        return A.bad("Invalid request.")
    if row["id"] == session["uid"]:
        return A.bad("You cannot change your own role.", 409)
    db.execute("UPDATE admins SET scope=%s WHERE id=%s", (scope, aid))
    A.audit("staff_scope_changed", f"{aid} -> {scope}")
    return jsonify(ok=True)


# =========================================================================== AUDIT LOG VIEWER - super only
def _audit_query(a, limit, offset=0):
    sql, args = " FROM audit_log WHERE 1=1", []
    q = a.get("q", "").strip()
    if q:
        sql += " AND (actor LIKE %s OR action LIKE %s OR detail LIKE %s)"
        args += [f"%{q}%"] * 3
    if a.get("actor", "").strip():
        sql += " AND actor=%s"
        args.append(a["actor"].strip())
    if a.get("action", "").strip():
        sql += " AND action=%s"
        args.append(a["action"].strip())
    for key, op in (("from", ">="), ("to", "<=")):
        v = _date(a.get(key, "")) if a.get(key) else None
        if v:
            sql += f" AND created_at {op} %s"
            args.append(v + (" 00:00:00" if key == "from" else " 23:59:59"))
    total = db.query_one("SELECT COUNT(*) AS n" + sql, tuple(args))["n"]
    rows = db.query("SELECT id, actor, action, detail, created_at" + sql + " ORDER BY id DESC LIMIT %s OFFSET %s",
                    tuple(args + [limit, offset]))
    return rows, total


@bp.route("/admin/audit")
@need("admin")
def admin_audit_page():
    return render_template("admin/audit.html", active="audit",
                           actions=[r["action"] for r in db.query("SELECT DISTINCT action FROM audit_log ORDER BY action")])


@bp.route("/admin/api/audit")
@need("admin")
def api_audit():
    page = max(1, _int(request.args.get("page")) or 1)
    per = 50
    rows, total = _audit_query(request.args, per, (page - 1) * per)
    return jsonify(ok=True, rows=rows, total=total, page=page, pages=max(1, -(-total // per)))


@bp.route("/admin/audit/export.csv")
@need("admin")
def audit_export():
    rows, _t = _audit_query(request.args, 5000)
    A.audit("audit_export", f"{len(rows)} rows")
    return A.send_csv("audit_log", ["ID", "Time", "Actor", "Action", "Detail"],
                      [[r["id"], r["created_at"], r["actor"], r["action"], r["detail"]] for r in rows])
